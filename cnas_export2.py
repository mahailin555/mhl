#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
CNAS授权项目导出（Firefox版）
安装依赖：
  pip3 install playwright beautifulsoup4 openpyxl lxml
  python3 -m playwright install firefox
然后直接运行：
  python3 cnas_export2.py
"""

import time
from pathlib import Path
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from playwright.sync_api import sync_playwright

# ── 参数 ──────────────────────────────────────────────────────────────
BASE_INFO_ID    = "e9e3a4bc1fa648a48f681db141fd28d4"
CERT_UPDATE_TS  = "2026-01-30"
VALIDATE        = "2030-11-18"
ATTACT_DATE     = "2026-01-30"
ORG_NAME        = "中联品检（佛山）检验技术有限公司"
CERT_NO         = "CNAS L1842"
OUTPUT_FILE     = str(Path.home() / "Desktop" / "中联品检佛山_CNAS授权项目.xlsx")

ABILITY_URL = (
    f"https://las.cnas.org.cn/LAS_FQ/publish/lab/checkLabObjListView.jsp"
    f"?baseInfoId={BASE_INFO_ID}&enstart=0&blueTooth=0&type=abilityL1"
    f"&orgEnOrCh=Ch&certUpdateTs={CERT_UPDATE_TS}&validate={VALIDATE}"
    f"&attactdate={ATTACT_DATE}&pageSize=100"
)

# ── 解析表格 ──────────────────────────────────────────────────────────
def parse_table(html: str) -> list[dict]:
    soup = BeautifulSoup(html, "lxml")
    items = []
    for tbl in soup.find_all("table"):
        rows = tbl.find_all("tr")
        if not rows:
            continue
        # 找表头（支持多行表头合并的情况）
        header_texts = []
        for r in rows[:3]:
            header_texts += [c.get_text(strip=True) for c in r.find_all(["th", "td"])]
        if not any("检测对象" in h for h in header_texts):
            continue

        # 建立列索引（遍历所有表头行）
        col = {}
        for ri, row in enumerate(rows[:3]):
            for ci, cell in enumerate(row.find_all(["th", "td"])):
                t = cell.get_text(strip=True)
                if "检测对象" in t and "obj" not in col:
                    col["obj"] = ci
                elif ("名称" in t or "项目" in t or "参数" in t) and "param" not in col:
                    col["param"] = ci
                elif ("标准" in t or "方法" in t) and "std" not in col:
                    col["std"] = ci

        if not col:
            continue

        print(f"  找到能力表，列映射: {col}，行数: {len(rows)}")

        # 跳过所有表头行（连续的 th 行）
        data_start = 0
        for ri, row in enumerate(rows):
            if row.find("th"):
                data_start = ri + 1
            else:
                break

        for row in rows[data_start:]:
            cells = [td.get_text(" ", strip=True) for td in row.find_all("td")]
            if len(cells) < 2:
                continue
            it = {
                "obj":   cells[col["obj"]]   if "obj"   in col and col["obj"]   < len(cells) else "",
                "param": cells[col["param"]] if "param" in col and col["param"] < len(cells) else "",
                "std":   cells[col["std"]]   if "std"   in col and col["std"]   < len(cells) else "",
            }
            if any(it.values()):
                items.append(it)
    return items


def get_total_pages(html: str) -> int:
    import re
    m = re.search(r"共\s*(\d+)\s*页", html)
    if m:
        return int(m.group(1))
    m = re.search(r"totalPage\s*[=:]\s*(\d+)", html)
    return int(m.group(1)) if m else 1


# ── 导出 Excel ────────────────────────────────────────────────────────
def export_excel(items: list[dict]) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CNAS授权项目"
    thin = Side(style="thin", color="BBBBBB")
    bdr  = Border(left=thin, right=thin, top=thin, bottom=thin)
    ctr  = Alignment(horizontal="center", vertical="center", wrap_text=True)
    lft  = Alignment(horizontal="left",   vertical="center", wrap_text=True)

    ws.merge_cells("A1:D1")
    ws["A1"] = f"机构：{ORG_NAME}    证书：{CERT_NO}    有效期至：{VALIDATE}"
    ws["A1"].font = Font(name="微软雅黑", bold=True, size=11)
    ws["A1"].alignment = lft
    ws.row_dimensions[1].height = 24

    for ci, (h, w) in enumerate(
        zip(["序号", "检测对象", "项目/参数", "检测标准（方法）"], [7, 25, 40, 55]), 1
    ):
        c = ws.cell(row=2, column=ci, value=h)
        c.font = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F6BB7")
        c.alignment = ctr
        c.border = bdr
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[2].height = 22

    stripe = PatternFill("solid", fgColor="EEF4FC")
    for ri, item in enumerate(items, 1):
        row = ri + 2
        for ci, val in enumerate([ri, item["obj"], item["param"], item["std"]], 1):
            c = ws.cell(row=row, column=ci, value=val)
            c.font = Font(name="微软雅黑", size=10)
            c.alignment = ctr if ci == 1 else lft
            c.border = bdr
            if ri % 2 == 0:
                c.fill = stripe
        ws.row_dimensions[row].height = 18

    ws.freeze_panes = "A3"
    wb.save(OUTPUT_FILE)
    return OUTPUT_FILE


# ── 主程序 ────────────────────────────────────────────────────────────
def main():
    print("=" * 55)
    print(f"机构：{ORG_NAME}")
    print(f"证书：{CERT_NO}    有效期至：{VALIDATE}")
    print("=" * 55)
    print("\n浏览器将自动打开，请稍等...\n")

    all_items: list[dict] = []

    with sync_playwright() as pw:
        # 使用 Firefox 避免 Google CDN 下载问题
        browser = pw.firefox.launch(headless=False)
        page = browser.new_page()
        page.set_extra_http_headers({"Accept-Language": "zh-CN,zh;q=0.9"})

        print(f"第1步：访问第1页...")
        page.goto(ABILITY_URL, wait_until="networkidle", timeout=60_000)
        time.sleep(2)

        html = page.content()
        total = get_total_pages(html)
        items = parse_table(html)
        all_items.extend(items)
        print(f"  第1页获取 {len(items)} 条，共 {total} 页")

        for pg in range(2, total + 1):
            print(f"第{pg}步：翻到第{pg}/{total}页...")
            # 尝试点击"下一页"按钮
            try:
                next_btn = page.locator("text=下一页").first
                if next_btn.is_visible(timeout=3_000):
                    next_btn.click()
                    page.wait_for_load_state("networkidle", timeout=20_000)
                    time.sleep(1.5)
                else:
                    raise Exception("未找到下一页按钮")
            except Exception:
                # 备用：直接跳转页码
                page_url = ABILITY_URL + f"&currentPage={pg}"
                page.goto(page_url, wait_until="networkidle", timeout=30_000)
                time.sleep(1.5)

            html = page.content()
            items = parse_table(html)
            all_items.extend(items)
            print(f"  获取 {len(items)} 条（累计 {len(all_items)} 条）")

        browser.close()

    if not all_items:
        print("\n未获取到数据，请截图发给我")
        return

    print(f"\n正在导出 Excel，共 {len(all_items)} 条...")
    out = export_excel(all_items)
    print("\n" + "=" * 55)
    print(f"完成！文件已保存到桌面：")
    print(f"  {out}")
    print("=" * 55)


if __name__ == "__main__":
    main()
