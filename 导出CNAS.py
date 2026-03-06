#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
直接运行即可，无需任何命令行参数。
需要安装：pip3 install requests beautifulsoup4 lxml openpyxl
"""

import time, sys
import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path

# ── 参数（直接从你提供的网址中提取）──────────────────────────────────
BASE_INFO_ID   = "e9e3a4bc1fa648a48f681db141fd28d4"
CERT_UPDATE_TS = "2026-01-30"
VALIDATE       = "2030-11-18"
ATTACT_DATE    = "2026-01-30"
ORG_NAME       = "中联品检（佛山）检验技术有限公司"
CERT_NO        = "CNAS L1842"
OUTPUT_FILE    = "中联品检佛山_CNAS授权项目.xlsx"

ABILITY_URL = (
    "https://las.cnas.org.cn/LAS_FQ/publish/lab/checkLabObjListView.jsp"
    f"?baseInfoId={BASE_INFO_ID}&enstart=0&blueTooth=0&type=abilityL1"
    f"&orgEnOrCh=Ch&certUpdateTs={CERT_UPDATE_TS}&validate={VALIDATE}"
    f"&attactdate={ATTACT_DATE}"
)

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/120.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "zh-CN,zh;q=0.9",
    "Referer": "https://las.cnas.org.cn/LAS_FQ/publish/externalQueryL1.jsp",
}

# ── 抓取函数 ──────────────────────────────────────────────────────────
def fetch_page(session, url):
    print(f"  正在访问: {url[:80]}...")
    for attempt in range(1, 4):
        try:
            r = session.get(url, timeout=30)
            r.encoding = r.apparent_encoding or "utf-8"
            return r.text
        except Exception as e:
            print(f"  第{attempt}次失败: {e}")
            if attempt < 3:
                time.sleep(3)
    raise RuntimeError(f"网络请求失败，请检查网络连接: {url}")

def parse_table(html):
    """解析检测能力表格，返回数据行列表"""
    soup = BeautifulSoup(html, "lxml")
    items = []

    # 找所有表格，寻找含有"检测对象"列的那个
    for table in soup.find_all("table"):
        rows = table.find_all("tr")
        if not rows:
            continue

        # 找表头行
        header_cells = rows[0].find_all(["th", "td"])
        headers = [c.get_text(strip=True) for c in header_cells]

        # 必须含"检测对象"才是目标表格
        if not any("检测对象" in h for h in headers):
            continue

        # 建立列索引
        col = {}
        for i, h in enumerate(headers):
            if "检测对象" in h:
                col["obj"] = i
            elif "项目" in h or "参数" in h:
                col["param"] = i
            elif "标准" in h or "方法" in h:
                col["std"] = i

        print(f"  找到能力表，共 {len(rows)-1} 行数据，列: {headers}")

        for row in rows[1:]:
            cells = [td.get_text(" ", strip=True) for td in row.find_all("td")]
            if len(cells) < 2:
                continue
            item = {
                "obj":   cells[col.get("obj",   0)] if col.get("obj")   is not None else "",
                "param": cells[col.get("param", 1)] if col.get("param") is not None else "",
                "std":   cells[col.get("std",   2)] if col.get("std")   is not None else "",
            }
            if any(item.values()):
                items.append(item)

    return items

def get_total_pages(html):
    """提取总页数"""
    import re
    m = re.search(r"共\s*(\d+)\s*页", html)
    if m:
        return int(m.group(1))
    m = re.search(r"totalPage\s*[=:]\s*(\d+)", html)
    return int(m.group(1)) if m else 1

# ── 导出 Excel ────────────────────────────────────────────────────────
def export_excel(items):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CNAS授权项目"

    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

    # 第1行：机构信息
    ws.merge_cells("A1:D1")
    ws["A1"] = f"机构：{ORG_NAME}    证书号：{CERT_NO}    有效期至：{VALIDATE}"
    ws["A1"].font = Font(name="微软雅黑", bold=True, size=11)
    ws["A1"].alignment = left
    ws.row_dimensions[1].height = 24

    # 第2行：表头
    headers = ["序号", "检测对象", "项目/参数", "检测标准（方法）"]
    widths  = [7, 25, 40, 55]
    for ci, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=2, column=ci, value=h)
        c.font      = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
        c.fill      = PatternFill("solid", fgColor="1F6BB7")
        c.alignment = center
        c.border    = border
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[2].height = 22

    # 数据行
    stripe = PatternFill("solid", fgColor="EEF4FC")
    for ri, item in enumerate(items, 1):
        row = ri + 2
        vals = [ri, item["obj"], item["param"], item["std"]]
        for ci, val in enumerate(vals, 1):
            c = ws.cell(row=row, column=ci, value=val)
            c.font      = Font(name="微软雅黑", size=10)
            c.alignment = center if ci == 1 else left
            c.border    = border
            if ri % 2 == 0:
                c.fill = stripe
        ws.row_dimensions[row].height = 18

    ws.freeze_panes = "A3"
    out = Path(OUTPUT_FILE)
    wb.save(out)
    return out.resolve()

# ── 主程序 ────────────────────────────────────────────────────────────
def main():
    print("=" * 50)
    print(f"目标机构：{ORG_NAME}")
    print(f"证书编号：{CERT_NO}")
    print("=" * 50)

    session = requests.Session()
    session.headers.update(HEADERS)

    # 先访问主页建立 Cookie
    print("\n第1步：建立会话...")
    try:
        session.get("https://las.cnas.org.cn/LAS_FQ/publish/externalQueryL1.jsp", timeout=15)
        time.sleep(1)
    except Exception as e:
        print(f"  （建立会话失败，继续尝试：{e}）")

    # 获取第1页
    print("\n第2步：获取授权检测项目列表...")
    html = fetch_page(session, ABILITY_URL)
    items = parse_table(html)
    total = get_total_pages(html)
    print(f"  共 {total} 页")

    # 翻页
    for pg in range(2, total + 1):
        print(f"\n第3步：翻页 {pg}/{total}...")
        page_url = ABILITY_URL + f"&currentPage={pg}"
        html = fetch_page(session, page_url)
        items.extend(parse_table(html))
        time.sleep(1)

    if not items:
        print("\n未获取到任何数据！")
        print("可能原因：网络不通 / 网站要求登录 / 页面结构变化")
        print(f"\n请用浏览器直接访问以下链接检查页面内容：\n{ABILITY_URL}")
        sys.exit(1)

    print(f"\n第4步：导出 Excel，共 {len(items)} 条项目...")
    out_path = export_excel(items)

    print("\n" + "=" * 50)
    print(f"完成！文件已保存到：")
    print(f"  {out_path}")
    print("=" * 50)

if __name__ == "__main__":
    main()
